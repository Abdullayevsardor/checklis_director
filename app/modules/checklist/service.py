"""
CHECKLIST SERVICE

Bu faylda checklist bilan bog‘liq asosiy database logikasi yoziladi.
Router faqat shu service funksiyalarini chaqiradi.
"""
import app.models  # MUHIM: barcha relationship modellarni yuklaydi

from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.modules.checklist.models import ChecklistSection


async def get_checklist_template(db: AsyncSession):
    """
    Aktiv checklist sectionlarini itemlari bilan olib keladi.

    Tartib:
    - avval section sort_order bo‘yicha
    - ichida itemlar sort_order bo‘yicha
    """

    query = (
        select(ChecklistSection)
        .where(ChecklistSection.is_active == True)
        .options(selectinload(ChecklistSection.items))
        .order_by(ChecklistSection.sort_order)
    )

    result = await db.execute(query)
    sections = result.scalars().unique().all()

    # Itemlarni ham faqat aktivlarini olib, tartiblab beramiz
    for section in sections:
        section.items = sorted(
            [item for item in section.items if item.is_active],
            key=lambda item: item.sort_order,
        )

    return sections




from io import BytesIO
from openpyxl import load_workbook
from sqlalchemy import select, delete
from fastapi import HTTPException

from app.modules.checklist.models import ChecklistSection, ChecklistItem
from app.modules.shift_checks.models import CheckAnswer, AnswerPhoto


def _parse_bool(value):
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    text = str(value).strip().lower()
    return text not in ("", "0", "false", "нет", "no", "n", "none")


def _normalize_header(value):
    if value is None:
        return ""
    return str(value).strip().lower()


def _has_header_row(headers):
    known_names = {
        "section_title",
        "section",
        "title_ru",
        "item_title",
        "item",
        "requires_photo",
        "item_sort",
        "sort_order",
        "item_id",
    }
    return any(_normalize_header(h) in known_names for h in headers)


async def clear_checklist(db: AsyncSession):
    """
    Barcha checklist itemlarini va sectionlarini o'chiradi.
    
    O'chirish tartibi muhim:
    1. Avval AnswerPhoto (answer_photos) - javoblarga bog'langan rasm URL'lari
    2. Keyin CheckAnswer (check_answers) - item javoblari
    3. Keyin ChecklistItem (checklist_items)
    4. Nihoyat ChecklistSection (checklist_sections)
    """
    await db.execute(delete(AnswerPhoto))
    await db.execute(delete(CheckAnswer))
    await db.execute(delete(ChecklistItem))
    await db.execute(delete(ChecklistSection))
    await db.commit()
    return {"message": "Checklist cleared successfully"}


async def add_checklist_from_excel(db, file, replace: bool = False):
    content = await file.read()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active

    if replace:
        # O'chirish tartibi muhim: CheckAnswer -> ChecklistItem -> ChecklistSection
        await db.execute(delete(CheckAnswer))
        await db.execute(delete(ChecklistItem))
        await db.execute(delete(ChecklistSection))

    default_section_title = "Общие вопросы"
    headers = [cell.value for cell in sheet[1]]
    has_header = _has_header_row(headers)
    normalized_headers = [_normalize_header(h) for h in headers]

    added_sections = 0
    added_items = 0
    section_cache = {}

    first_row = 2 if has_header else 1
    for row in sheet.iter_rows(min_row=first_row, values_only=True):
        row = list(row)
        if not row or all(cell is None for cell in row):
            continue

        if has_header:
            row_data = {
                normalized_headers[i]: row[i] if i < len(row) else None
                for i in range(len(normalized_headers))
            }
            section_title = (
                row_data.get("section_title")
                or row_data.get("section")
                or row_data.get("title_ru")
                or default_section_title
            )
            section_sort = row_data.get("section_sort") or row_data.get("sort_order") or 0
            item_title = row_data.get("item_title") or row_data.get("title") or row[0]
            requires_photo = row_data.get("requires_photo")
            item_sort = row_data.get("item_sort") or row_data.get("sort_order") or 0
        else:
            section_title = default_section_title
            section_sort = 0
            item_title = row[0] if len(row) > 0 else None
            requires_photo = row[1] if len(row) > 1 else False
            item_sort = row[2] if len(row) > 2 else 0

        if not item_title:
            continue

        section_title = str(section_title).strip() or default_section_title
        item_title = str(item_title).strip()

        section = section_cache.get(section_title)
        if not section:
            section = await db.scalar(
                select(ChecklistSection).where(
                    ChecklistSection.title_ru == section_title
                )
            )
            if not section:
                section = ChecklistSection(
                    title_ru=section_title,
                    sort_order=int(section_sort or 0),
                )
                db.add(section)
                await db.flush()
                added_sections += 1
            section_cache[section_title] = section

        item = ChecklistItem(
            section_id=section.id,
            title_ru=item_title,
            requires_photo=_parse_bool(requires_photo),
            sort_order=int(item_sort or 0),
        )

        db.add(item)
        added_items += 1

    await db.commit()

    return {
        "message": "Checklist added from Excel",
        "added_sections": added_sections,
        "added_items": added_items,
    }

async def update_checklist_from_excel(db, file):
    content = await file.read()
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    has_header = _has_header_row(headers)
    normalized_headers = [_normalize_header(h) for h in headers]

    updated_items = 0
    not_found_items = []

    first_row = 2 if has_header else 1
    for row in sheet.iter_rows(min_row=first_row, values_only=True):
        row = list(row)
        if not row or all(cell is None for cell in row):
            continue

        if has_header:
            row_data = {
                normalized_headers[i]: row[i] if i < len(row) else None
                for i in range(len(normalized_headers))
            }
            item_id = row_data.get("item_id") or row_data.get("id")
            item_title = row_data.get("item_title") or row_data.get("title")
            requires_photo = row_data.get("requires_photo")
            item_sort = row_data.get("item_sort") or row_data.get("sort_order")
        else:
            item_id = row[0] if len(row) > 0 else None
            item_title = row[1] if len(row) > 1 else None
            requires_photo = row[2] if len(row) > 2 else None
            item_sort = row[3] if len(row) > 3 else None

        if not item_id:
            continue

        item = await db.scalar(
            select(ChecklistItem).where(ChecklistItem.id == int(item_id))
        )

        if not item:
            not_found_items.append(item_id)
            continue

        if item_title:
            item.title_ru = str(item_title).strip()

        if requires_photo is not None:
            item.requires_photo = _parse_bool(requires_photo)

        if item_sort is not None:
            item.sort_order = int(item_sort)

        updated_items += 1

    await db.commit()

    return {
        "message": "Checklist updated from Excel",
        "updated_items": updated_items,
        "not_found_items": not_found_items,
    }